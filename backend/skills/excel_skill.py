"""Excel 表格处理 Skill：读取、创建和编辑 .xlsx 文件。"""

import logging
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook


logger = logging.getLogger("agent.skills.excel")


def _normalize(value):
    """把单元格值转成可 JSON 序列化的基础类型。"""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    return value


def excel_list_sheets(file_path: str):
    """列出 Excel 文件中的所有工作表名称。"""
    path = Path(file_path).expanduser().resolve()
    workbook = load_workbook(path, read_only=True)
    try:
        sheets = workbook.sheetnames
    finally:
        workbook.close()
    logger.info("Excel 工作表列表：file=%s，sheets=%s", path, sheets)
    return {"file": str(path), "sheets": sheets}


def excel_read_sheet(
    file_path: str,
    sheet_name: str = "",
    max_rows: int = 50,
    max_cols: int = 20,
):
    """读取 Excel 指定工作表，默认读取第一个工作表。"""
    path = Path(file_path).expanduser().resolve()
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        if sheet_name and sheet_name not in workbook.sheetnames:
            raise ValueError(f"工作表不存在: {sheet_name}")
        worksheet = workbook[sheet_name] if sheet_name else workbook[workbook.sheetnames[0]]
        rows = []
        for row in worksheet.iter_rows(
            min_row=1,
            max_row=max(1, max_rows),
            max_col=max(1, max_cols),
            values_only=True,
        ):
            rows.append([_normalize(cell) for cell in row])
    finally:
        workbook.close()

    logger.info(
        "Excel 读取完成：file=%s，sheet=%s，rows=%d",
        path,
        worksheet.title,
        len(rows),
    )
    return {
        "file": str(path),
        "sheet": worksheet.title,
        "row_count": len(rows),
        "rows": rows,
    }


def excel_write_cell(
    file_path: str,
    sheet_name: str = "Sheet1",
    row: int = 1,
    column: int = 1,
    value: str = "",
):
    """写入 Excel 指定单元格；文件不存在时自动创建。"""
    path = Path(file_path).expanduser().resolve()
    if path.exists():
        workbook = load_workbook(path)
    else:
        workbook = Workbook()

    try:
        if sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]
        else:
            worksheet = workbook.active
            worksheet.title = sheet_name
        worksheet.cell(row=max(1, int(row)), column=max(1, int(column)), value=value)
        workbook.save(path)
        logger.info(
            "Excel 写入单元格：file=%s，sheet=%s，row=%s，column=%s",
            path,
            worksheet.title,
            row,
            column,
        )
        return {
            "file": str(path),
            "sheet": worksheet.title,
            "row": int(row),
            "column": int(column),
            "value": str(value),
        }
    finally:
        workbook.close()


def excel_create_file(
    file_path: str,
    sheet_name: str = "Sheet1",
    headers: list = None,
    rows: list = None,
):
    """创建新的 Excel 文件，支持表头和二维数据。"""
    path = Path(file_path).expanduser().resolve()
    workbook = Workbook()
    try:
        worksheet = workbook.active
        worksheet.title = sheet_name or "Sheet1"

        if headers:
            worksheet.append([str(header) for header in headers])

        for row in rows or []:
            if isinstance(row, dict) and headers:
                worksheet.append([str(row.get(header, "")) for header in headers])
            else:
                worksheet.append([_normalize(cell) for cell in row])

        workbook.save(path)
        logger.info("Excel 文件创建完成：file=%s，sheet=%s", path, worksheet.title)
        return {
            "file": str(path),
            "sheet": worksheet.title,
            "headers": headers or [],
            "row_count": len(rows or []),
        }
    finally:
        workbook.close()


EXCEL_SKILL = {
    "name": "excel",
    "description": "读取、创建和编辑 Excel 表格文件（.xlsx）。",
    "instructions": "处理 Excel 时，先查看工作表结构，再读取或写入数据。",
    "tool_definitions": [
        {
            "type": "function",
            "function": {
                "name": "excel_list_sheets",
                "description": "列出 Excel 文件中的所有工作表名称。",
                "parameters": {
                    "type": "object",
                    "properties": {
                        "file_path": {"type": "string", "description": "Excel 文件路径"},
                    },
                    "required": ["file_path"],
                },
            },
        },
        {
            "type": "function",
            "function": {
                "name": "excel_read_sheet",
                "description": "读取 Excel 指定工作表的数据。",
                "parameters": {
                    "type": "object",
                    "properties": {
                        "file_path": {"type": "string", "description": "Excel 文件路径"},
                        "sheet_name": {"type": "string", "description": "工作表名称，默认第一个"},
                        "max_rows": {"type": "integer", "description": "最多读取行数，默认 50"},
                        "max_cols": {"type": "integer", "description": "最多读取列数，默认 20"},
                    },
                    "required": ["file_path"],
                },
            },
        },
        {
            "type": "function",
            "function": {
                "name": "excel_write_cell",
                "description": "写入 Excel 指定单元格，文件不存在时自动创建。",
                "parameters": {
                    "type": "object",
                    "properties": {
                        "file_path": {"type": "string", "description": "Excel 文件路径"},
                        "sheet_name": {"type": "string", "description": "工作表名称，默认 Sheet1"},
                        "row": {"type": "integer", "description": "行号，从 1 开始"},
                        "column": {"type": "integer", "description": "列号，从 1 开始"},
                        "value": {"type": "string", "description": "要写入的值"},
                    },
                    "required": ["file_path", "row", "column", "value"],
                },
            },
        },
        {
            "type": "function",
            "function": {
                "name": "excel_create_file",
                "description": "创建新的 Excel 文件，支持表头和二维数据。",
                "parameters": {
                    "type": "object",
                    "properties": {
                        "file_path": {"type": "string", "description": "要创建的 Excel 文件路径"},
                        "sheet_name": {"type": "string", "description": "工作表名称，默认 Sheet1"},
                        "headers": {
                            "type": "array",
                            "items": {"type": "string"},
                            "description": "表头列表",
                        },
                        "rows": {
                            "type": "array",
                            "description": "数据行，支持二维数组或对象数组",
                        },
                    },
                    "required": ["file_path"],
                },
            },
        },
    ],
    "tool_functions": {
        "excel_list_sheets": excel_list_sheets,
        "excel_read_sheet": excel_read_sheet,
        "excel_write_cell": excel_write_cell,
        "excel_create_file": excel_create_file,
    },
}
