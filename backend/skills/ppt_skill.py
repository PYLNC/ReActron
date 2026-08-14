"""PPT Skill：读取演示文稿结构、提取文本和创建 .pptx 文件。"""

import logging
from pathlib import Path

from pptx import Presentation


logger = logging.getLogger("agent.skills.ppt")


def _slide_title(slide) -> str:
    title = slide.shapes.title
    return title.text if title is not None else ""


def ppt_list_slides(file_path: str):
    """列出 PPT 文件中的幻灯片数量和标题。"""
    path = Path(file_path).expanduser().resolve()
    presentation = Presentation(path)
    slides = [
        {"index": index + 1, "title": _slide_title(slide)}
        for index, slide in enumerate(presentation.slides)
    ]
    logger.info("PPT 幻灯片列表：file=%s，count=%d", path, len(slides))
    return {"file": str(path), "slide_count": len(slides), "slides": slides}


def ppt_read_text(file_path: str):
    """提取 PPT 中每张幻灯片的标题和正文文本。"""
    path = Path(file_path).expanduser().resolve()
    presentation = Presentation(path)
    slides = []
    for index, slide in enumerate(presentation.slides):
        texts = []
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for paragraph in shape.text_frame.paragraphs:
                text = "".join(run.text for run in paragraph.runs) or paragraph.text
                if text.strip():
                    texts.append(text.strip())
        slides.append({
            "index": index + 1,
            "title": _slide_title(slide),
            "texts": texts,
        })
    logger.info("PPT 文本提取完成：file=%s，slides=%d", path, len(slides))
    return {"file": str(path), "slide_count": len(slides), "slides": slides}


def _fill_content_slide(slide, title: str, bullets: list):
    slide.shapes.title.text = title or ""
    body = None
    for shape in slide.placeholders:
        if shape.placeholder_format.idx == 1:
            body = shape
            break
    if body is None:
        return

    text_frame = body.text_frame
    text_frame.clear()
    for index, bullet in enumerate(bullets or []):
        paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
        paragraph.text = str(bullet)


def ppt_create_file(file_path: str, title: str = "", slides: list = None):
    """创建新的 PPT 文件，slides 为包含 title 和 bullets 的对象数组。"""
    path = Path(file_path).expanduser().resolve()
    presentation = Presentation()

    if title:
        title_slide = presentation.slides.add_slide(presentation.slide_layouts[0])
        title_slide.shapes.title.text = title

    for item in slides or []:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        _fill_content_slide(slide, item.get("title", ""), item.get("bullets", []))

    presentation.save(path)
    logger.info("PPT 文件创建完成：file=%s，slides=%d", path, len(slides or []))
    return {
        "file": str(path),
        "title": title,
        "slide_count": len(presentation.slides),
    }


PPT_SKILL = {
    "name": "ppt",
    "description": "读取 PPT 结构、提取文本和创建 .pptx 演示文稿。",
    "instructions": "创建 PPT 时，slides 参数使用 [{\"title\": \"标题\", \"bullets\": [\"要点\"]}] 结构。",
    "tool_definitions": [
        {
            "type": "function",
            "function": {
                "name": "ppt_list_slides",
                "description": "列出 PPT 文件中的幻灯片数量和标题。",
                "parameters": {
                    "type": "object",
                    "properties": {
                        "file_path": {"type": "string", "description": "PPT 文件路径"},
                    },
                    "required": ["file_path"],
                },
            },
        },
        {
            "type": "function",
            "function": {
                "name": "ppt_read_text",
                "description": "提取 PPT 中每张幻灯片的标题和正文文本。",
                "parameters": {
                    "type": "object",
                    "properties": {
                        "file_path": {"type": "string", "description": "PPT 文件路径"},
                    },
                    "required": ["file_path"],
                },
            },
        },
        {
            "type": "function",
            "function": {
                "name": "ppt_create_file",
                "description": "创建新的 PPT 文件。",
                "parameters": {
                    "type": "object",
                    "properties": {
                        "file_path": {"type": "string", "description": "要创建的 PPT 文件路径"},
                        "title": {"type": "string", "description": "封面标题"},
                        "slides": {
                            "type": "array",
                            "items": {
                                "type": "object",
                                "properties": {
                                    "title": {"type": "string"},
                                    "bullets": {
                                        "type": "array",
                                        "items": {"type": "string"},
                                    },
                                },
                            },
                            "description": "幻灯片列表",
                        },
                    },
                    "required": ["file_path"],
                },
            },
        },
    ],
    "tool_functions": {
        "ppt_list_slides": ppt_list_slides,
        "ppt_read_text": ppt_read_text,
        "ppt_create_file": ppt_create_file,
    },
}
