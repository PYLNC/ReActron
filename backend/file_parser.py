"""上传文件解析：支持 Excel、PPT、CSV、文本和 JSON。"""

import csv
import json
from dataclasses import dataclass
from pathlib import Path
from pathlib import PurePosixPath
from typing import Any
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from pptx import Presentation


class DocumentExtractionError(ValueError):
    """Safe, fixed-code extraction failure."""

    code = "document_parse_failed"


class DocumentLimitExceeded(DocumentExtractionError):
    code = "document_limit_exceeded"


class UnsupportedDocumentType(DocumentExtractionError):
    code = "unsupported_document_type"


class DocumentNoText(DocumentExtractionError):
    code = "document_no_text"


class DocumentExtractionTimeout(DocumentExtractionError):
    code = "document_extraction_timeout"


@dataclass(frozen=True)
class DocumentSegment:
    text: str
    section: str
    locator: dict[str, Any]


def _check_zip_limits(path: Path, *, required: tuple[str, ...] = ()) -> None:
    """解析库打开 ZIP 前，拒绝明显的展开量与压缩比炸弹。"""
    try:
        with ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > 2_000:
                raise DocumentLimitExceeded("ZIP entry limit exceeded")
            total = 0
            names = set()
            for item in entries:
                name = item.filename.replace("\\", "/")
                parts = PurePosixPath(name).parts
                is_symlink = (item.external_attr >> 16) & 0o170000 == 0o120000
                if (
                    not name or "\x00" in name or name.startswith("/") or
                    ".." in parts or (parts and ":" in parts[0]) or
                    item.flag_bits & 1 or is_symlink or name in names
                ):
                    raise DocumentExtractionError("Invalid ZIP container")
                names.add(name)
                total += item.file_size
                if (
                    item.file_size > 25 * 1024 * 1024
                    or total > 100 * 1024 * 1024
                    or item.file_size > max(item.compress_size, 1) * 100
                ):
                    raise DocumentLimitExceeded("ZIP expansion limit exceeded")
            if any(name not in names for name in required):
                raise DocumentExtractionError("Invalid ZIP container")
    except BadZipFile as exc:
        raise DocumentExtractionError("Invalid ZIP container") from exc


def _parse_excel(path: Path) -> dict:
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        rows = []
        for row in sheet.iter_rows(max_row=20, values_only=True):
            rows.append([str(cell) if cell is not None else "" for cell in row])
        return {
            "type": "excel",
            "sheets": workbook.sheetnames,
            "sheet": sheet.title,
            "preview_rows": rows,
        }
    finally:
        workbook.close()


def _parse_pptx(path: Path) -> dict:
    presentation = Presentation(path)
    slides = []
    for index, slide in enumerate(presentation.slides):
        texts = []
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for paragraph in shape.text_frame.paragraphs:
                text = "".join(run.text for run in paragraph.runs) or paragraph.text
                if text.strip():
                    texts.append(text.strip())
        slides.append({"index": index + 1, "texts": texts[:10]})
    return {"type": "pptx", "slide_count": len(slides), "slides": slides[:10]}


def _parse_csv(path: Path) -> dict:
    with path.open("r", encoding="utf-8-sig", errors="replace") as handle:
        reader = csv.reader(handle)
        rows = [row for _, row in zip(range(20), reader)]
    return {"type": "csv", "preview_rows": rows}


def _parse_text(path: Path) -> dict:
    text = path.read_text(encoding="utf-8", errors="replace")
    return {
        "type": "text",
        "preview": text[:2000],
        "char_count": len(text),
    }


def parse_uploaded_file(path: Path) -> dict:
    """根据扩展名解析文件并返回摘要。"""
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        _check_zip_limits(path, required=("[Content_Types].xml", "xl/workbook.xml"))
        return _parse_excel(path)
    if suffix == ".pptx":
        _check_zip_limits(path, required=("[Content_Types].xml", "ppt/presentation.xml"))
        return _parse_pptx(path)
    if suffix == ".csv":
        return _parse_csv(path)
    if suffix in (".txt", ".md", ".json", ".log"):
        return _parse_text(path)
    return {"type": "unsupported", "preview": ""}


def extract_document(path: Path, *, timeout_seconds: float = 20.0) -> list[DocumentSegment]:
    """Extract bounded full text in a killable child process."""
    from backend.rag_extract import extract_document as extract
    return extract(path, timeout_seconds=timeout_seconds)


def to_json_safe(data) -> str:
    return json.dumps(data, ensure_ascii=False, default=str)
